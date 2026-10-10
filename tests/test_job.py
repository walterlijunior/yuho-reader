# Yuho Reader - baixa demonstrações contábeis do EDINET e monta planilhas Excel.
# Copyright (C) 2026 walterlijunior
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

import threading
from datetime import date

import pytest
from openpyxl import load_workbook

from core.api import Cancelled
from core.companies import Company, parse_codelist
from core.job import JobOptions, company_folder, list_library, run_job
from tests.fake_edinet import TODAY, make_codelist_zip


@pytest.fixture
def capcom():
    return parse_codelist(make_codelist_zip())[0]


def opcoes(capcom, tmp_path, anos=(2025, 2024), **kwargs):
    return JobOptions(company=capcom, fiscal_year_ends=[date(a, 3, 31) for a in anos],
                      output_dir=tmp_path / "saida", **kwargs)


@pytest.mark.parametrize("nome_en, nome_jp, esperado", [
    ("CAPCOM CO., LTD.", "株式会社カプコン", "E02417_CAPCOM_CO_LTD"),
    ("  Seven & i Holdings Co., Ltd.  ", "", "E02417_Seven_i_Holdings_Co_Ltd"),
    ('A/B: "C" <D>|E?*.', "", "E02417_A_B_C_D_E"),
    ("", "株式会社カプコン", "E02417_株式会社カプコン"),
    ("...", "", "E02417"),
    ("", "", "E02417"),
])
def test_nome_da_pasta_sem_espacos_virgulas_ou_ponto_final(tmp_path, nome_en, nome_jp, esperado):
    empresa = Company(edinet_code="E02417", name_en=nome_en, name_jp=nome_jp)
    pasta = company_folder(tmp_path, empresa)
    assert pasta == tmp_path / esperado
    assert not any(ch in pasta.name for ch in " ,.")


def test_tarefa_completa_gera_arquivos_e_planilha(client, capcom, tmp_path):
    progresso, registro = [], []
    resultado = run_job(
        client, opcoes(capcom, tmp_path, semiannual=True, amendments=True, xbrl=True, pdf=True),
        on_progress=lambda pct, status: progresso.append(pct), on_log=registro.append, today=TODAY,
    )
    pasta = tmp_path / "saida" / "E02417_CAPCOM_CO_LTD"
    assert resultado.folder == pasta
    nomes = sorted(p.name for p in pasta.iterdir())
    assert "Anual_2025-03_S100AN25_csv.zip" in nomes
    assert "Anual_2025-03_S100AN25_xbrl.zip" in nomes
    assert "Anual_2025-03_S100AN25.pdf" in nomes
    assert "Anual_2024-03_retif_2024-07-03_S100RT24_csv.zip" in nomes
    assert resultado.downloaded == 15 and resultado.skipped == 0
    assert progresso[0] == 0 and progresso[-1] == 100 and progresso == sorted(progresso)

    wb = load_workbook(resultado.workbook)
    assert wb.sheetnames == [
        "INDICE", "PAINEL", "DEMONSTRACOES",
        "Anual_2025-03", "Semestral_2024-09", "Anual_2024-03",
        "Anual_2024-03_retif_2024-07-03", "Semestral_2023-09",
    ]

    indice = wb["INDICE"]
    assert [c.value for c in indice[2]][:3] == ["Anual_2025-03", "Anual", "S100AN25"]
    assert indice["B5"].value == "Anual (retificação)"
    assert "Anual_2025-03_S100AN25.pdf" in indice["H2"].value
    link = "https://disclosure2dl.edinet-fsa.go.jp/searchdocument/pdf/S100AN25.pdf"
    assert indice["I2"].value == link and indice["I2"].hyperlink.target == link

    painel = wb["PAINEL"]
    assert [c.value for c in painel[1]] == [
        "Taxonomia", "ID do elemento", "Conta (japonês)", "Fluxo/Saldo",
        "Anual_2025-03", "Semestral_2024-09", "Anual_2024-03",
        "Anual_2024-03_retif_2024-07-03", "Semestral_2023-09",
    ]
    linhas = {linha[1]: linha for linha in painel.iter_rows(min_row=2, values_only=True)}
    assert list(linhas) == ["jppfs_cor:Assets", "jppfs_cor:NetSales", "jppfs_cor:OperatingIncome"]
    assert linhas["jppfs_cor:NetSales"] == ("J-GAAP", "jppfs_cor:NetSales", "売上高", "fluxo",
                                            169604, 56400, 152410, 152411, 74900)
    assert linhas["jppfs_cor:Assets"][3] == "saldo"

    # Formatação igual em todas as abas.
    for aba in wb.worksheets:
        topo = aba["A1"]
        assert topo.font.bold and topo.font.color.rgb.endswith("76B900")
        assert topo.fill.fgColor.rgb.endswith("1A1A1A")
        assert aba.freeze_panes in ("A2", "E2")
        assert aba.auto_filter.ref
    assert painel["E2"].number_format == "#,##0;[Red]-#,##0"

    demonstracoes = wb["DEMONSTRACOES"]
    assert demonstracoes.max_row == 1 + 5 * 3
    assert [c.value for c in demonstracoes[1]][:4] == ["Relatório", "Taxonomia", "Fluxo/Saldo", "要素ID"]

    aba = wb["Anual_2025-03"]
    assert aba.max_row == 11 and aba["J1"].value == "valor_numerico"
    texto = aba["I11"]
    assert texto.value == "=texto que parece fórmula" and texto.data_type == "s"

    assert any("Planilha pronta" in m for m in registro)


def test_nao_baixa_de_novo_o_que_ja_existe(make_client, fake, capcom, tmp_path):
    run_job(make_client(), opcoes(capcom, tmp_path), today=TODAY)
    downloads = fake.count("/documents/")
    resultado = run_job(make_client(), opcoes(capcom, tmp_path), today=TODAY)
    assert fake.count("/documents/") == downloads
    assert resultado.downloaded == 0 and resultado.skipped == 2


def test_segunda_busca_acrescenta_a_planilha(make_client, capcom, tmp_path):
    run_job(make_client(), opcoes(capcom, tmp_path, anos=[2025]), today=TODAY)
    resultado = run_job(make_client(), opcoes(capcom, tmp_path, anos=[2024]), today=TODAY)
    painel = load_workbook(resultado.workbook)["PAINEL"]
    assert [c.value for c in painel[1]][4:] == ["Anual_2025-03", "Anual_2024-03"]


def test_nada_encontrado(client, capcom, tmp_path):
    resultado = run_job(client, opcoes(capcom, tmp_path, anos=[2020]), today=TODAY)
    assert resultado.workbook is None
    assert resultado.missing == ["Anual_2020-03"]
    assert not resultado.folder.exists()


def test_cancelar_no_meio(make_client, capcom, tmp_path):
    cancelar = threading.Event()
    client = make_client(cancel_event=cancelar)

    def ao_progredir(pct, status):
        if pct >= 10:
            cancelar.set()

    with pytest.raises(Cancelled):
        run_job(client, opcoes(capcom, tmp_path), on_progress=ao_progredir, today=TODAY)
    assert not company_folder(tmp_path / "saida", capcom).exists()


def test_biblioteca(client, capcom, tmp_path):
    assert list_library(tmp_path / "saida") == []
    run_job(client, opcoes(capcom, tmp_path, semiannual=True), today=TODAY)
    (item,) = list_library(tmp_path / "saida")
    assert (item.edinet_code, item.sec_code, item.name) == ("E02417", "9697", "CAPCOM CO., LTD.")
    assert item.years == [2023, 2024, 2025]
    assert item.workbook is not None and item.workbook.is_file()
