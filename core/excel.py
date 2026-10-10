# Yuho Reader - baixa demonstrações contábeis do EDINET e monta planilhas Excel.
# Copyright (C) 2026 walterlijunior
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""Monta a planilha Excel de uma empresa: INDICE, PAINEL, DEMONSTRACOES e uma aba por relatório."""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .api import pdf_link
from .parser import COLUMNS, NUMERIC, current_values, extract_statements
from .search import FoundDoc

HEADER_FILL = PatternFill("solid", fgColor="1A1A1A")
HEADER_FONT = Font(name="Segoe UI", bold=True, color="76B900")
LINK_FONT = Font(color="0563C1", underline="single")
NUMBER_FORMAT = "#,##0;[Red]-#,##0"
MAX_CELL = 32767  # limite de caracteres de uma célula do Excel
MAX_SHEET_NAME = 31


@dataclass
class Report:
    doc: FoundDoc
    files: list[str] = field(default_factory=list)
    data: pd.DataFrame | None = None  # None quando o relatório não tem CSV


def _clean(value: object) -> object:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL]
    return value


def _append(ws, row: list) -> None:
    ws.append([_clean(value) for value in row])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            cell.number_format = NUMBER_FORMAT
        elif isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"  # texto do EDINET, não fórmula


def _text_width(value: object) -> int:
    if isinstance(value, (int, float)):
        return len(f"{value:,.0f}")
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in str(value or ""))


def _finish(ws) -> None:
    """Cabeçalho escuro com texto verde, linha congelada, filtro e larguras ajustadas."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    sample = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300), values_only=True))
    for index in range(ws.max_column):
        width = max((_text_width(row[index]) for row in sample), default=0)
        ws.column_dimensions[get_column_letter(index + 1)].width = max(10, min(width + 3, 60))


def _sheet_name(label: str, used: set[str]) -> str:
    name = label
    if len(name) > MAX_SHEET_NAME:
        name = name.replace("Semestral_", "Sem_").replace("Anual_", "An_")
    name = name[:MAX_SHEET_NAME]
    counter = 2
    while name in used:
        suffix = f"~{counter}"
        name = name[: MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    used.add(name)
    return name


def build_workbook(path: Path, reports: list[Report]) -> Path:
    """Grava a planilha e devolve o caminho usado.

    Se o arquivo estiver aberto no Excel, grava com outro nome em vez de falhar.
    """
    reports = sorted(reports, key=lambda report: report.doc.sort_key)
    with_data = [report for report in reports if report.data is not None]
    statements = {report.doc.label: extract_statements(report.data) for report in with_data}

    workbook = Workbook()

    # ---------------------------------------------------------------- INDICE
    ws = workbook.active
    ws.title = "INDICE"
    ws.append([
        "Rótulo", "Tipo", "docID", "Descrição", "Início do período", "Fim do período",
        "Data de envio", "Arquivos salvos", "Link do PDF",
    ])
    for report in reports:
        doc = report.doc
        kind = doc.kind + (" (retificação)" if doc.amendment else "")
        link = pdf_link(doc.doc_id)
        _append(ws, [
            doc.label, kind, doc.doc_id, doc.description, doc.period_start, doc.period_end,
            doc.submitted, ", ".join(report.files), link,
        ])
        cell = ws.cell(row=ws.max_row, column=9)
        cell.hyperlink = link
        cell.font = LINK_FONT
    _finish(ws)

    # ---------------------------------------------------------------- PAINEL
    ws = workbook.create_sheet("PAINEL")
    labels = [report.doc.label for report in with_data]
    ws.append(["Taxonomia", "ID do elemento", "Conta (japonês)", "Fluxo/Saldo", *labels])
    values = {label: current_values(frame) for label, frame in statements.items()}
    accounts: dict[tuple[str, str], dict] = {}
    for label in labels:
        for key, row in values[label].items():
            accounts.setdefault(key, row)
    for key, row in accounts.items():
        _append(ws, [
            row["taxonomia"], row["要素ID"], row["項目名"], row["natureza"],
            *[values[label].get(key, {}).get(NUMERIC) for label in labels],
        ])
    _finish(ws)
    ws.freeze_panes = "E2"

    # --------------------------------------------------------- DEMONSTRACOES
    ws = workbook.create_sheet("DEMONSTRACOES")
    ws.append(["Relatório", "Taxonomia", "Fluxo/Saldo", *COLUMNS, NUMERIC])
    for label in labels:
        for row in statements[label].to_dict("records"):
            _append(ws, [label, row["taxonomia"], row["natureza"], *[row[c] for c in COLUMNS], row[NUMERIC]])
    _finish(ws)

    # ----------------------------------------------------- uma aba por relatório
    used = {"INDICE", "PAINEL", "DEMONSTRACOES"}
    for report in with_data:
        ws = workbook.create_sheet(_sheet_name(report.doc.label, used))
        ws.append([*COLUMNS, NUMERIC])
        for row in report.data.to_dict("records"):
            _append(ws, [*[row[c] for c in COLUMNS], row[NUMERIC]])
        _finish(ws)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(path)
    except PermissionError:
        path = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d-%H%M%S}{path.suffix}")
        workbook.save(path)
    return path
